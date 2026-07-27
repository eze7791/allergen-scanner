"""Import products and menu items from the curated allergen spreadsheet.

Reads two sheets from an .xlsx workbook:
  - "Main Products Allergens": raw supplier products -> FoodItem rows.
  - "Allergens Matrix New": menu dishes grouped by category headers -> Recipe rows,
    with allergens set directly (not composed from products).

Usage:
    python import_menu.py <xlsx_path> --restaurant-name "BrewDog" --admin-pin <pin>
"""
from __future__ import annotations

import argparse
import asyncio
import re

import openpyxl
from sqlalchemy import delete, select

from database import async_session
from models import (
    AllergenORM,
    FoodItemAllergenORM,
    FoodItemORM,
    RecipeAllergenORM,
    RecipeORM,
    RestaurantORM,
    UserORM,
    UserRole,
)
from auth import generate_join_code, hash_pin
from seed_allergens import seed_allergens

PRODUCTS_SHEET = "Main Products Allergens"
MATRIX_SHEET = "Allergens Matrix New"

# Spreadsheet column header -> the legal allergen name seeded by seed_allergens.py
COLUMN_TO_ALLERGEN = {
    "GLUTEN": "Cereals containing gluten",
    "CRUSTACEANS": "Crustaceans",
    "EGGS": "Eggs",
    "FISH": "Fish",
    "PEANUTS": "Peanuts",
    "SOYA": "Soybeans",
    "MILK": "Milk",
    "NUTS": "Nuts",
    "CELERY": "Celery",
    "MUSTARD": "Mustard",
    "SESAME": "Sesame seeds",
    "SULPHITES": "Sulphur dioxide and sulphites",
    "LUPIN": "Lupin",
    "MOLLUSCS": "Molluscs",
}
ALLERGEN_COLUMNS = list(COLUMN_TO_ALLERGEN.keys())  # sheet column order, matches header row

CELL_RE = re.compile(r"^(Y|MC|N)(?:\((.*)\))?$", re.IGNORECASE)


def parse_cell(value) -> tuple[str | None, str | None]:
    """Returns (kind, note) where kind is 'Y', 'MC', or None (nothing to record)."""
    if value is None:
        return None, None
    text = str(value).strip()
    if not text:
        return None, None
    m = CELL_RE.match(text.upper())
    if not m:
        print(f"  ! unrecognized cell value {value!r}, skipping")
        return None, None
    kind = m.group(1)
    note = m.group(2)
    if kind == "N":
        return None, None
    return kind, note


def parse_row_allergens(row: tuple) -> list[tuple[str, bool, str | None]]:
    """row = (name, gluten, crustaceans, ..., molluscs). Returns [(allergen_name, certain, note), ...]."""
    results = []
    for col_name, cell in zip(ALLERGEN_COLUMNS, row[1:15]):
        kind, note = parse_cell(cell)
        if kind is None:
            continue
        results.append((COLUMN_TO_ALLERGEN[col_name], kind == "Y", note))
    return results


def load_products(path: str) -> list[tuple[str, list[tuple[str, bool, str | None]]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[PRODUCTS_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    products = []
    for row in rows[1:]:  # skip header
        name = row[0]
        if not name or not str(name).strip():
            continue
        allergens = parse_row_allergens(row)
        products.append((str(name).strip(), allergens))
    return products


def load_dishes(path: str) -> list[tuple[str, str, list[tuple[str, bool, str | None]]]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[MATRIX_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    dishes = []
    current_category = None
    for row in rows:
        name = row[0]
        no_allergen_cells = all(c is None for c in row[1:15])
        # Most category headers repeat the column labels ('GLUTEN', 'CRUSTACEANS', ...)
        # in place of allergen values, but a couple (e.g. "SALADS") are bare - fall back
        # to "all-caps name, no allergen data, no parenthetical" to catch those too.
        # Footnote rows ("SUGAR STRAND*(...)", "All Pizza May Contain (...)") always have
        # a parenthetical or aren't all-caps, so they're excluded from this fallback.
        is_header_row = row[1] == "GLUTEN" or (
            name and no_allergen_cells and str(name).strip().isupper() and "(" not in str(name)
        )
        if is_header_row:
            if name and str(name).strip().upper() != "MENU ITEM":
                current_category = str(name).strip()
            continue
        if not name or not str(name).strip():
            continue
        allergens = parse_row_allergens(row)
        if not allergens:
            # Spacer/footnote rows ("All Pizza May Contain...", "SUGAR STRAND*...",
            # or dishes with no allergen data entered yet) - nothing to import.
            continue
        dishes.append((str(name).strip(), current_category, allergens))
    return dishes


async def get_or_create_restaurant(db, name: str, admin_pin: str) -> RestaurantORM:
    result = await db.execute(select(RestaurantORM).where(RestaurantORM.name == name))
    restaurant = result.scalar_one_or_none()
    if restaurant:
        # Re-running the import resets the admin PIN to the one provided, so
        # re-imports always leave you with a known-good admin credential.
        restaurant.admin_pin_hash = hash_pin(admin_pin)
        await db.commit()
        return restaurant

    restaurant = RestaurantORM(name=name, join_code=generate_join_code(), admin_pin_hash=hash_pin(admin_pin))
    db.add(restaurant)
    await db.commit()
    await db.refresh(restaurant)

    admin = UserORM(restaurant_id=restaurant.id, role=UserRole.admin)
    db.add(admin)
    await db.commit()
    return restaurant


async def import_menu(path: str, restaurant_name: str, admin_pin: str):
    async with async_session() as db:
        restaurant = await get_or_create_restaurant(db, restaurant_name, admin_pin)
        await seed_allergens(restaurant.id)

        allergen_result = await db.execute(
            select(AllergenORM).where(AllergenORM.restaurant_id == restaurant.id)
        )
        allergen_by_name = {a.name: a.id for a in allergen_result.scalars().all()}

        # Idempotent re-import: wipe existing products/dishes (and their allergen
        # links) for this tenant.
        await db.execute(
            delete(FoodItemAllergenORM).where(
                FoodItemAllergenORM.food_item_id.in_(
                    select(FoodItemORM.id).where(FoodItemORM.restaurant_id == restaurant.id)
                )
            )
        )
        await db.execute(
            delete(RecipeAllergenORM).where(
                RecipeAllergenORM.recipe_id.in_(
                    select(RecipeORM.id).where(RecipeORM.restaurant_id == restaurant.id)
                )
            )
        )
        await db.execute(delete(FoodItemORM).where(FoodItemORM.restaurant_id == restaurant.id))
        await db.execute(delete(RecipeORM).where(RecipeORM.restaurant_id == restaurant.id))
        await db.commit()

        products = load_products(path)
        for name, allergens in products:
            item = FoodItemORM(restaurant_id=restaurant.id, name=name)
            item.allergen_links = [
                FoodItemAllergenORM(allergen_id=allergen_by_name[allergen_name], certain=certain, note=note)
                for allergen_name, certain, note in allergens
                if allergen_name in allergen_by_name
            ]
            db.add(item)

        dishes = load_dishes(path)
        for name, category, allergens in dishes:
            recipe = RecipeORM(restaurant_id=restaurant.id, name=name, category=category)
            recipe.allergen_links = [
                RecipeAllergenORM(allergen_id=allergen_by_name[allergen_name], certain=certain, note=note)
                for allergen_name, certain, note in allergens
                if allergen_name in allergen_by_name
            ]
            db.add(recipe)

        await db.commit()

        print(f"Restaurant: {restaurant.name} (id={restaurant.id})")
        print(f"  join_code: {restaurant.join_code}")
        print(f"  admin_pin: {admin_pin}")
        print(f"Imported {len(products)} products, {len(dishes)} dishes.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--restaurant-name", default="BrewDog")
    parser.add_argument("--admin-pin", required=True)
    args = parser.parse_args()
    asyncio.run(import_menu(args.xlsx_path, args.restaurant_name, args.admin_pin))
